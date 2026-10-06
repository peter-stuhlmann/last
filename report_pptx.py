"""
PowerPoint-Bericht für den LLM-Härtetest – in einfacher Sprache.

Durchgehendes Bild: Der Server ist ein Geschichtenerzähler. Jede Anfrage ist ein Kind,
das ihm eine Frage stellt. Er antwortet Wort für Wort. Fachbegriffe stehen jeweils klein
darunter, damit auch Technik-Leute alles wiederfinden.
"""

from collections import Counter

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from llm_stresstest import (
    PROBE_KINDS,
    capacity,
    economics,
    fmt,
    meets_slo,
    rate_ok,
    recommended_concurrency,
    round_users,
    verdict,
)

# ── Farben ──────────────────────────────────────────────────────────────────
INK, INK2, MUTED = "111827", "374151", "6B7280"
DARK_BG, CARD, PANEL = "0F172A", "F1F5F9", "EEF4FB"
S1, S2, S3, S4 = "2A78D6", "EB6834", "1BAF7A", "4A3AA7"
GOOD, WARN, BAD, INFO = "16A06D", "D99A00", "D63D3C", "6B7280"
GRID, REF = "E5E7EB", "9CA3AF"
STATUS = {"good": GOOD, "warn": WARN, "bad": BAD, "info": INFO}
W = 13.333

rgb = RGBColor.from_string


# ════════════════════════════════════════════════════════════════════════════
#  Alltags-Übersetzungen
# ════════════════════════════════════════════════════════════════════════════

def words(tokens):
    """Tokens → ungefähre Wörter (1 Wort ≈ 1,3 Tokens)."""
    return tokens / 1.33 if tokens is not None else None


def pages(tokens):
    """Tokens → ungefähre Buchseiten (1 Seite ≈ 300 Wörter ≈ 400 Tokens)."""
    return tokens / 400 if tokens is not None else None


def time_feel(s):
    if s is None:
        return "–"
    if s < 0.3:
        return "schneller als ein Wimpernschlag"
    if s < 1:
        return "fast sofort"
    if s < 2.5:
        return "kurz – so lange wie einmal Luft holen"
    if s < 6:
        return "merklich – man wartet ein paar Atemzüge"
    if s < 20:
        return "lang – man wird ungeduldig"
    return "sehr lang – viele würden aufgeben"


def speed_feel(tps):
    """Tokens/s pro Person. Lesen: ca. 4–5 Wörter/s ≈ 6 Tokens/s."""
    if tps is None:
        return "–"
    if tps >= 25:
        return "viel schneller, als man lesen kann"
    if tps >= 8:
        return "schneller, als man lesen kann"
    if tps >= 4:
        return "etwa so schnell, wie man liest"
    return "langsamer, als man liest – zäh"


def stall_feel(s):
    """Längste Pause mitten in einer Antwort."""
    if s is None:
        return "–"
    if s < 0.3:
        return "fällt nicht auf"
    if s < 1:
        return "kaum spürbar"
    if s < 3:
        return "man merkt kurze Hänger"
    return "deutliche Hänger – wirkt wie eingefroren"


def stall_status(s):
    return None if s is None else GOOD if s < 1 else WARN if s < 3 else BAD


def cap(t):
    return t[:1].upper() + t[1:] if t else t


def mins(m):
    """'1 Minute', '3 Minuten', '1,5 Minuten'."""
    m = float(m)
    return fmt(m, 0 if m == int(m) else 1) + (" Minute" if m == 1 else " Minuten")


def n0(v):
    return fmt(v, 0)


def secs(v, d=None):
    if v is None:
        return "–"
    if d is None:
        d = 2 if v < 1 else 1 if v < 10 else 0
    return fmt(v, d, " s")


def approx(v):
    return "–" if v is None else n0(round_users(v))


def per_min(rate):
    return fmt(rate * 60, 0 if rate * 60 >= 10 else 1) if rate is not None else "–"


def secs_pair(a, b):
    """Zwei Zeiten mit so vielen Nachkommastellen, dass der Unterschied sichtbar wird."""
    d = 3 if max(a or 0, b or 0) < 1 else 2 if max(a or 0, b or 0) < 10 else 1
    return secs(a, d), secs(b, d)


def ctx_label(tok):
    k = f"{round(tok / 1000)}k" if tok >= 1000 else str(tok)
    p = pages(tok)
    return f"{k} (≈{fmt(p, 0 if p >= 2 else 1)} S.)"


# ════════════════════════════════════════════════════════════════════════════
#  Bausteine
# ════════════════════════════════════════════════════════════════════════════

class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, dark=False):
        sl = self.prs.slides.add_slide(self.blank)
        f = sl.background.fill
        f.solid()
        f.fore_color.rgb = rgb(DARK_BG if dark else "FFFFFF")
        return sl

    def save(self, path):
        self.prs.save(path)


def runs(p, parts, size, color, font="Calibri"):
    """parts: Liste aus str oder (str, {bold, color, italic})."""
    for part in parts:
        txt, opt = (part, {}) if isinstance(part, str) else part
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(opt.get("size", size))
        r.font.bold = opt.get("bold", False)
        r.font.italic = opt.get("italic", False)
        r.font.name = opt.get("font", font)
        r.font.color.rgb = rgb(opt.get("color", color))


def text(sl, x, y, w, h, content, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font="Calibri", italic=False, space_after=0):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        parts = para if isinstance(para, list) else [para]
        parts = [(pp, {"bold": bold, "italic": italic, "font": font}) if isinstance(pp, str) else pp
                 for pp in parts]
        runs(p, parts, size, color, font)
    return tb


def rect(sl, x, y, w, h, color, rounded=True, radius=0.06):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if rounded:
        shp.adjustments[0] = radius
    return shp


def dot(sl, x, y, d, color, label=None, label_color="FFFFFF", size=12):
    shp = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if label is not None:
        text(sl, x, y, d, d, str(label), size=size, bold=True, color=label_color,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return shp


def title(sl, t, sub=None):
    text(sl, 0.6, 0.42, W - 1.2, 0.7, t, size=30, bold=True, font="Cambria")
    if sub:
        text(sl, 0.6, 1.1, W - 1.2, 0.45, sub, size=16, color=INK2)


def footer(sl, s):
    text(sl, 0.6, 6.95, W - 1.2, 0.35, s, size=10.5, color=MUTED, italic=True)


def notes(sl, s):
    sl.notes_slide.notes_text_frame.text = s


def explain(sl, how, meaning, x=8.95, y=1.7, w=3.85, h=5.0):
    """Erklärkasten rechts: 'So liest du das Bild' + 'Was das bedeutet'."""
    rect(sl, x, y, w, h, PANEL)
    text(sl, x + 0.3, y + 0.25, w - 0.6, 0.35, "So liest du das Bild", size=14, bold=True, color=S1)
    text(sl, x + 0.3, y + 0.62, w - 0.6, 1.9, how, size=12.5, color=INK2)
    text(sl, x + 0.3, y + 2.55, w - 0.6, 0.35, "Was das bedeutet", size=14, bold=True, color=S1)
    text(sl, x + 0.3, y + 2.92, w - 0.6, h - 3.1, meaning, size=12.5, color=INK)


def stat(sl, x, y, w, h, value, label, meaning=None, color=INK, fill=CARD, vsize=32):
    rect(sl, x, y, w, h, fill)
    text(sl, x + 0.25, y + 0.18, w - 0.5, 0.62, value, size=vsize, bold=True, color=color,
         anchor=MSO_ANCHOR.MIDDLE)
    text(sl, x + 0.25, y + 0.86, w - 0.5, 0.3, label, size=12.5, bold=True, color=INK2)
    if meaning:
        text(sl, x + 0.25, y + 1.16, w - 0.5, h - 1.2, meaning, size=11.5, color=MUTED)


# ── Diagramme ───────────────────────────────────────────────────────────────

def style_chart(ch, ttl, y_title=None, x_title=None, legend=True):
    ch.font.size = Pt(11)
    ch.font.name = "Calibri"
    ch.font.color.rgb = rgb(INK2)
    ch.has_title = bool(ttl)
    if ttl:
        ch.chart_title.text_frame.text = ttl
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = rgb(INK)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(GRID)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(10)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(REF)
    ca.tick_labels.font.size = Pt(10)
    ca.tick_labels.font.color.rgb = rgb(MUTED)
    try:
        ca.has_major_gridlines = False
    except Exception:
        pass
    for ax, t in ((va, y_title), (ca, x_title)):
        if t:
            ax.has_title = True
            ax.axis_title.text_frame.text = t
            r = ax.axis_title.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(10.5)
            r.font.bold = False
            r.font.color.rgb = rgb(MUTED)


def style_series(ch, colors, dashed=(), markers=True):
    for i, ser in enumerate(ch.plots[0].series):
        c = colors[i % len(colors)]
        ser.smooth = False
        ser.format.line.color.rgb = rgb(c)
        ser.format.line.width = Pt(1.5 if i in dashed else 2.5)
        if i in dashed:
            ser.format.line.dash_style = MSO_LINE.DASH
            ser.marker.style = XL_MARKER_STYLE.NONE
        elif markers:
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 8
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(c)
            ser.marker.format.line.color.rgb = rgb("FFFFFF")
        else:
            ser.marker.style = XL_MARKER_STYLE.NONE


def line_chart(sl, x, y, w, h, cats, series, ttl=None, y_title=None, x_title=None,
               colors=(S1, S2, S3, S4), dashed=(), number_format="0.0"):
    cd = CategoryChartData(number_format=number_format)
    cd.categories = [str(c) for c in cats]
    for name, vals in series:
        cd.add_series(name, vals)
    ch = sl.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w),
                             Inches(h), cd).chart
    style_chart(ch, ttl, y_title, x_title, legend=len(series) > 1)
    style_series(ch, colors, dashed)
    ch.value_axis.minimum_scale = 0
    return ch


def xy_chart(sl, x, y, w, h, series, ttl=None, y_title=None, x_title="Minuten seit Start",
             colors=(S1, S2, S3, S4)):
    cd = XyChartData()
    for name, pts in series:
        s = cd.add_series(name)
        for px, py in pts:
            s.add_data_point(px, py)
    ch = sl.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y),
                             Inches(w), Inches(h), cd).chart
    style_chart(ch, ttl, y_title, x_title, legend=len(series) > 1)
    style_series(ch, colors, markers=False)
    return ch


def bar_chart(sl, x, y, w, h, cats, vals, ttl=None, y_title=None, x_title=None, color=S2,
              number_format="0.0"):
    cd = CategoryChartData(number_format=number_format)
    cd.categories = [str(c) for c in cats]
    cd.add_series(ttl or "Werte", vals)
    ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w),
                             Inches(h), cd).chart
    style_chart(ch, ttl, y_title, x_title, legend=False)
    plot = ch.plots[0]
    plot.gap_width = 40
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(color)
    return ch


def bars_chart(sl, x, y, w, h, cats, series, y_title=None, colors=(S1, S2), number_format="0",
               label_format=None):
    """Säulendiagramm mit mehreren Reihen und Werten über den Säulen."""
    cd = CategoryChartData(number_format=number_format)
    cd.categories = [str(c) for c in cats]
    for name, vals in series:
        cd.add_series(name, vals)
    ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w),
                             Inches(h), cd).chart
    style_chart(ch, None, y_title, None, legend=len(series) > 1)
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    for i, ser in enumerate(plot.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
    if label_format:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = label_format
        dl.number_format_is_linked = False
        dl.font.size = Pt(10)
        dl.font.color.rgb = rgb(INK2)
    ch.value_axis.minimum_scale = 0
    return ch


def eur(v, d=0):
    return fmt(float(v), d, " €") if v is not None else "–"


def table(sl, x, y, w, rows, col_w=None, font=11, row_h=0.34):
    nr, nc = len(rows), len(rows[0])
    tbl = sl.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = tbl.cell(r, c)
            val, color = rows[r][c] if isinstance(rows[r][c], tuple) else (rows[r][c], None)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(font)
            run.font.name = "Calibri"
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(DARK_BG)
                run.font.bold = True
                run.font.color.rgb = rgb("FFFFFF")
            else:
                cell.fill.fore_color.rgb = rgb("FFFFFF" if r % 2 else "F8FAFC")
                run.font.color.rgb = rgb(color or INK)
                run.font.bold = bool(color)
    return tbl


# ════════════════════════════════════════════════════════════════════════════
#  Der Bericht
# ════════════════════════════════════════════════════════════════════════════

def headline_cards(summary):
    """Die vier Zahlen für die Entscheidung (Wert, Beschriftung, Erläuterung, Farbe) –
    für die Empfehlungsfolie und den PDF-Bericht."""
    cfg = summary.get("config", {})
    cap_, eco, pr = capacity(summary), economics(summary), summary.get("probes")
    users, think = cap_.get("users"), cap_.get("think_time", 60)
    target = cfg.get("target_users") or 0
    cards = [(f"≈ {approx(users)}" if users else "0", "aktive Nutzer gleichzeitig",
              (f"gebraucht: {target}" if target else f"bei {n0(think)} s Pause zwischen zwei Fragen"),
              (GOOD if not target or (users or 0) >= target else WARN if (users or 0) >= 0.7 * target else BAD)
              if users else BAD)]
    if eco:
        be = eco.get("breakeven_pct")
        cards.append((eur(eco["own_month"]), "Kosten pro Monat",
                      f"eigener Server; die Cloud würde bei voller Nutzung {eur(eco['cloud_full_month'])} kosten",
                      INK))
        cards.append(((f"{fmt(be, 0)} %" if be <= 100 else "nie") if be is not None else "–",
                       "günstiger als die Cloud ab",
                       "so viel Auslastung während der Arbeitszeit", GOOD if be is not None and be <= 60
                       else WARN if be is not None and be <= 100 else BAD))
    if pr and pr.get("load_pct") is not None:
        cards.append((f"{fmt(pr['load_pct'], 0)} %", "richtige Antworten unter Last",
                      f"ohne Last: {fmt(pr['alone_pct'], 0)} %" if pr.get("alone_pct") is not None else "",
                      GOOD if pr["load_pct"] >= (pr.get("alone_pct") or 100) - 3 else WARN))
    return cards


def result_rows(summary):
    """Alle Ergebnisse als Ampel-Zeilen (Farbe, Bereich, Ergebnis) – für die Zusammenfassungsfolie
    und den einseitigen PDF-Bericht, damit beide dasselbe sagen."""
    cfg = summary.get("config", {})
    slo_t, slo_u = cfg.get("slo_ttft", 2.0), cfg.get("slo_user_tps", 10.0)
    base, ctx, spike, soak = (summary.get(k) for k in ("baseline", "context", "spike", "soak"))
    ramp = summary.get("ramp") or {}
    steps = sorted(ramp.get("steps", []), key=lambda x: x["concurrency"])
    brk, saturated = ramp.get("break"), ramp.get("saturated")
    rate = summary.get("rate") or {}
    rstages, rbrk = rate.get("stages", []), rate.get("break")
    rec = recommended_concurrency(summary)
    rec_step = next((x for x in steps if x["concurrency"] == rec), None)
    cap_, eco, pr = capacity(summary), economics(summary), summary.get("probes")
    rows = []  # (Farbe, Bereich, Ergebnis)
    if base:
        tps, tt = base["user_tps_p50"], base["ttft_p50"]
        rows.append((GOOD if tps and tps >= slo_u else WARN, "Tempo für eine Person",
                     f"{fmt(tps, 0)} Tokens/s – {speed_feel(tps)}"))
        rows.append((GOOD if tt is not None and tt <= slo_t else WARN, "Wartezeit ohne Andrang",
                     f"{secs(tt)} bis zum ersten Wort – {time_feel(tt)}"
                     + (f" (Antwort nach Nachdenken: {secs(base['ttfc_p50'])})"
                        if base.get("ttfc_p50") is not None else "")))
    if steps:
        rows.append((GOOD if rec else BAD, "Gleichzeitig flüssig",
                     f"bis {rec} Fragen gleichzeitig" if rec else "keine Stufe erreicht die Ziele"))
        if brk:
            rows.append((BAD, "Belastungsgrenze", f"ab {brk['concurrency']} gleichzeitig: {brk['reason']}"))
        elif saturated:
            rows.append((WARN, "Belastungsgrenze", f"ab {saturated} gleichzeitig keine Mehrleistung"))
        else:
            rows.append((GOOD, "Belastungsgrenze", f"bis {steps[-1]['concurrency']} gleichzeitig "
                                                   "nicht erreicht"))
    if rstages:
        r_ok = cap_.get("rate_ok")
        rows.append((GOOD if r_ok else BAD, "Echter Andrang",
                     (f"bis {per_min(r_ok)} Fragen/min flüssig" if r_ok else "schon die kleinste Menge zu viel")
                     + (f", ab {per_min(rbrk['rate'])}/min wächst die Schlange" if rbrk else "")))
    stall = (rec_step or {}).get("stall_p95")
    if stall is not None:
        rows.append((stall_status(stall), "Ruckeln in der Antwort",
                     f"höchstens {secs(stall)} Pause am Stück – {stall_feel(stall)}"))
    pre = [s for s in steps if s.get("preemptions") is not None]
    if pre:
        bad_ = [s for s in pre if s["preemptions"] and (not rec or s["concurrency"] <= rec)]
        rows.append((WARN if bad_ else GOOD, "Notizblock (KV-Cache)",
                     f"läuft schon bei {bad_[0]['concurrency']} gleichzeitig über" if bad_ else
                     "reicht im empfohlenen Bereich"))
    pc = (base or {}).get("prefix_compare")
    if pc and pc.get("ttft_unique"):
        rows.append((INFO, "Gemeinsamer Einleitungstext",
                     "erstes Wort nach {} statt {}".format(*secs_pair(pc["ttft_shared"], pc["ttft_unique"]))))
    if ctx and ctx.get("max_ok"):
        lim = ctx.get("limit")
        rows.append((WARN if lim and lim["mode"] != "einzeln" else GOOD, "Lange Texte",
                     f"bis ≈ {n0(pages(ctx['max_ok']))} Buchseiten"
                     + (f"; mehrere gleichzeitig bis ≈ {n0(pages(lim['input_len']))} Seiten"
                        if lim and lim["mode"] != "einzeln" else "")))
    if spike:
        ok_ = spike.get("recovered") and spike["burst"]["errors"] == 0
        rows.append((GOOD if ok_ else BAD, "Großer Ansturm",
                     f"{spike['size']} auf einmal – " + ("gut überstanden" if ok_ else "nicht sauber überstanden")))
    if pr and pr.get("load_pct") is not None:
        ref = pr.get("alone_pct") if pr.get("alone_pct") is not None else 100
        rows.append((GOOD if ref - pr["load_pct"] <= 3 else WARN if ref - pr["load_pct"] <= 10 else BAD,
                     "Antworten unter Last", f"{fmt(pr['load_pct'], 0)} % richtig (allein {fmt(ref, 0)} %)"))
    if eco and eco.get("breakeven_pct") is not None:
        be = eco["breakeven_pct"]
        rows.append((GOOD if be <= 60 else WARN if be <= 100 else BAD, "Kosten",
                     f"{eur(eco['own_month'])}/Monat – " + (f"günstiger als Cloud ab {fmt(be, 0)} % Auslastung"
                                                            if be <= 100 else "Cloud wäre günstiger")))
    dg = ((soak or {}).get("degradation") or {}).get("tps_change_pct")
    if dg is not None:
        rows.append((GOOD if dg > -10 else BAD, "Ausdauer",
                     f"nach {mins(soak['minutes'])} {'unverändert' if abs(dg) < 2 else fmt(dg, 0) + ' %'}"))
    return rows


def build_report(summary, path):
    d = Deck()
    meta = summary.get("meta", {})
    cfg = summary.get("config", {})
    model = cfg.get("model") or "?"
    slo_t = cfg.get("slo_ttft", 2.0)
    slo_u = cfg.get("slo_user_tps", 10.0)
    base = summary.get("baseline")
    ramp = summary.get("ramp") or {}
    steps = sorted(ramp.get("steps", []), key=lambda s: s["concurrency"])
    brk = ramp.get("break")
    ctx = summary.get("context")
    spike = summary.get("spike")
    soak = summary.get("soak")
    ctl = summary.get("control")
    quality = summary.get("quality") or {"overall": "–", "items": []}
    rec = recommended_concurrency(summary)
    cap_ = capacity(summary)
    users, think = cap_.get("users"), cap_.get("think_time", 60)
    rate = summary.get("rate") or {}
    rstages = sorted(rate.get("stages", []), key=lambda s: s["rate"])
    rbrk = rate.get("break")
    saturated = ramp.get("saturated")
    rec_step = next((s for s in steps if s["concurrency"] == rec), None)
    eco = economics(summary)
    vd = verdict(summary)
    pr = summary.get("probes")
    target = cfg.get("target_users") or 0
    top = max(steps, key=lambda s: s["out_tps"] or 0) if steps else None
    sat = None
    if top:
        for s in steps:
            if (s["out_tps"] or 0) >= 0.9 * top["out_tps"]:
                sat = s["concurrency"]
                break
    in_len, out_len = cfg.get("input_len", 1024), cfg.get("output_len", 256)
    q_color = {"hoch": GOOD, "mittel": WARN, "niedrig": BAD}.get(quality["overall"], INFO)

    # ═══ 1 Titel ═══
    sl = d.slide(dark=True)
    text(sl, 0.8, 1.55, 11.5, 0.45, "BELASTUNGSTEST", size=15, color="94A3B8", bold=True)
    text(sl, 0.8, 2.0, 11.8, 1.0, "Wie belastbar ist unser KI-Server?", size=40, color="FFFFFF",
         bold=True, font="Cambria")
    text(sl, 0.8, 3.05, 11.8, 0.6, model, size=24, color="CBD5E1")
    started = (meta.get("started") or "")[:16].replace("T", " ")
    text(sl, 0.8, 4.4, 11.5, 1.2, [f"Getestet am {started}",
                                   f"Server: {cfg.get('base_url', '')}  ·  Messrechner: {meta.get('host', '–')}"],
         size=15, color="94A3B8")
    for i, c in enumerate((S1, S2, S3)):
        dot(sl, 0.8 + i * 0.35, 6.3, 0.2, c)
    notes(sl, "Automatisch erzeugt von llm_stresstest.py. Rohdaten: summary.json und "
              "raw_requests.jsonl im Ergebnisordner.")

    # ═══ 1b Empfehlung für die Entscheidung ═══
    sl = d.slide()
    title(sl, "Unsere Empfehlung")
    vcol = STATUS[vd["status"]]
    rect(sl, 0.6, 1.3, W - 1.2, 1.1, vcol)
    text(sl, 1.0, 1.3, W - 2.0, 1.1, vd["headline"], size=28, bold=True, color="FFFFFF",
         anchor=MSO_ANCHOR.MIDDLE, font="Cambria")
    cards = headline_cards(summary)
    cw = (W - 1.2 - 0.25 * (len(cards) - 1)) / len(cards)
    for i, (v, l, m, col) in enumerate(cards):
        stat(sl, 0.6 + i * (cw + 0.25), 2.65, cw, 1.95, v, l, m, color=col, vsize=30)
    yy = 4.85
    text(sl, 0.6, yy, 6, 0.35, "Warum?", size=15, bold=True, color=S1)
    for i, (st_, t_) in enumerate(vd["reasons"][:5]):
        dot(sl, 0.65, yy + 0.47 + i * 0.36, 0.18, STATUS[st_])
        text(sl, 0.98, yy + 0.4 + i * 0.36, W - 1.6, 0.34, t_, size=13, color=INK)
    footer(sl, "Kosten = Abschreibung + Strom, Annahmen auf der Folie „Was kostet uns das?“. "
               "Nutzerzahl nach dem Gesetz von Little – eine Schätzung, keine Garantie.")
    notes(sl, "Urteil: grün, wenn Kapazität (gegen den Bedarf, falls angegeben), Antwortqualität unter "
              "Last und Stabilität passen; gelb bei knapper Kapazität, leicht sinkender Qualität, "
              "Cloud günstiger oder unsauberer Messung; rot bei zu wenig Kapazität, deutlich "
              "schlechteren Antworten oder Absturz.")

    # ═══ 2 Kurzfassung ═══
    sl = d.slide()
    title(sl, "Das Wichtigste auf einen Blick")
    cards = []
    if base:
        cards.append((fmt(base["user_tps_p50"], 1), "Tokens pro Sekunde für eine Person",
                      f"≈ {n0(words(base['user_tps_p50']))} Wörter/s – {speed_feel(base['user_tps_p50'])}", S1))
        cards.append((secs(base["ttft_p50"]), "bis die Antwort beginnt (ohne Andrang)",
                      cap(time_feel(base["ttft_p50"])), INK))
    if users:
        cards.append((f"≈ {approx(users)}", "aktive Nutzer – ohne dass es zäh wird",
                      f"bei {n0(think)} s Denkpause zwischen zwei Fragen"
                      + (f" · {rec} Fragen gleichzeitig in Arbeit" if rec else ""), GOOD))
    elif steps:
        cards.append(("0", "aktive Nutzer – ohne dass es zäh wird",
                      f"Ziel: Antwort beginnt nach ≤ {fmt(slo_t, 1)} s, Tempo ≥ {n0(slo_u)} Tokens/s",
                      BAD))
    if steps:
        if brk:
            cards.append((str(brk["concurrency"]), "Fragen gleichzeitig – ab hier geht es kaputt",
                          f"Grund: {brk['reason']}", BAD))
        elif saturated:
            cards.append((str(saturated), "Fragen gleichzeitig – mehr bringt nichts",
                          "Die Leistung stieg nicht mehr, die Ziele waren verfehlt.", WARN))
        else:
            cards.append((f"> {steps[-1]['concurrency']}", "Fragen gleichzeitig – noch nicht kaputt",
                          "Die Grenze liegt höher als getestet.", GOOD))
    if top:
        cards.append((n0(top["out_tps"]), "Tokens/s für alle zusammen (maximal)",
                      f"≈ {n0(words(top['out_tps']) * 60 / 300)} Buchseiten pro Minute", INK))
    if ctx and ctx.get("max_ok"):
        cards.append((f"≈ {n0(pages(ctx['max_ok']))} Seiten", "längster Text, den er lesen konnte",
                      f"{n0(ctx['max_ok'])} Tokens", INK))
    for i, (v, l, m, col) in enumerate(cards[:6]):
        r, c = divmod(i, 3)
        stat(sl, 0.6 + c * 4.1, 1.45 + r * 2.1, 3.85, 1.85, v, l, m, color=col)
    y = 5.8
    rect(sl, 0.6, y, W - 1.2, 0.9, CARD)
    dot(sl, 0.9, y + 0.29, 0.32, q_color)
    text(sl, 1.4, y + 0.14, 5.5, 0.4, f"Verlässlichkeit der Messung: {quality['overall']}", size=17,
         bold=True)
    goods = sum(1 for i in quality["items"] if i["status"] == "good")
    text(sl, 1.4, y + 0.5, 11, 0.35, f"{goods} von {len(quality['items'])} Prüfungen ohne "
         "Auffälligkeiten – Details auf der Folie „Wie sehr können wir den Zahlen trauen?“",
         size=12, color=MUTED)
    notes(sl, "Tokens sind Wortstücke; 1 Wort ≈ 1,3 Tokens. 'Empfohlen' = höchste getestete "
              f"Parallelität ohne Fehler mit TTFT p95 ≤ {slo_t} s und medianer Decode-Rate ≥ {slo_u} tok/s. "
              "Aktive Nutzer = Gesetz von Little: Fragen pro Sekunde × (Antwortdauer + Denkpause), "
              + ("aus dem Test mit zufälligen Ankünften (höchste Rate, die die Ziele einhält)."
                 if cap_.get("source") == "rate" else "aus der empfohlenen Stufe der Ramp."))

    # ═══ 3 So haben wir getestet ═══
    sl = d.slide()
    title(sl, "So haben wir getestet")
    rect(sl, 0.6, 1.45, 5.6, 5.25, PANEL)
    text(sl, 0.95, 1.75, 5.0, 0.4, "Stell dir vor …", size=18, bold=True, color=S1)
    text(sl, 0.95, 2.3, 4.95, 4.2, [
        "… der Server ist ein Geschichtenerzähler.",
        "Jede Frage an ihn ist wie ein Kind, das ihn etwas fragt. Er antwortet Wort für Wort.",
        "Wir wollten wissen: Wie schnell ist er? Wie viele Kinder schafft er gleichzeitig? "
        "Und wann wird es ihm zu viel?",
        "Dafür haben wir ihm echte Fragen geschickt – erst eine, dann immer mehr – und mit der "
        "Stoppuhr gemessen."], size=15, color=INK, space_after=10)
    tests = [("Einer allein", "Wie schnell ist er, wenn nur ein Kind fragt?"),
             ("Immer mehr gleichzeitig", "2, 4, 8 … Kinder fragen gleichzeitig. Wann wird es zäh?")]
    if rstages:
        tests.append(("Echter Andrang", "Kinder kommen zufällig, wie im Alltag. Wann bildet sich "
                                        "eine Schlange?"))
    tests += [("Lange Texte", "Wie viele Buchseiten kann er lesen, bevor er antwortet?"),
              ("Großer Ansturm", "Die ganze Klasse kommt auf einmal. Hält er das aus?"),
              ("Ausdauer", "Viele Minuten Dauerbetrieb. Wird er müde?"),
              ("Kontrolle", "Zum Schluss noch einmal wie am Anfang. Ist er gleich schnell?")]
    colors = (S1, S2, "B45309", S3, S4, "0E7490", MUTED) if rstages else \
             (S1, S2, S3, S4, "0E7490", MUTED)
    gap = min(0.86, 5.3 / len(tests))
    for i, (h, b) in enumerate(tests):
        yy = 1.5 + i * gap
        dot(sl, 6.7, yy + 0.05, 0.5, colors[i], i + 1, size=15)
        text(sl, 7.4, yy, 5.4, 0.35, h, size=16, bold=True)
        text(sl, 7.4, yy + 0.36, 5.4, 0.4, b, size=13, color=MUTED)
    notes(sl, "Baseline, Concurrency-Ramp (zeitbasiert, eingeschwungenes Messfenster, Feinsuche, "
              "Sättigungs-Abbruch), offene Last mit Poisson-Ankünften, Kontext-Sweep, Spike mit "
              "Erholungsprüfung, Soak, Kontrollmessung. "
              f"Längen: {'gemischt (log-normal)' if cfg.get('length_mix') == 'mixed' else 'fest'}; "
              f"gemeinsamer System-Prompt: {cfg.get('shared_prefix') or 0} Tokens; "
              f"Seed: {cfg.get('seed', '–')}.")

    # ═══ 4 Wörterbuch ═══
    sl = d.slide()
    title(sl, "Kleines Wörterbuch", "Die wichtigsten Begriffe – einmal erklärt, dann benutzt")
    terms = [("Token", "Ein Wortstück. Ein Wort besteht meist aus 1–2 Tokens. 100 Tokens ≈ 75 Wörter.",
              "Token"),
             ("Wartezeit", "Wie lange es dauert, bis das erste Wort der Antwort kommt.",
              "Time to First Token (TTFT)"),
             ("Tempo", "Wie viele Tokens pro Sekunde eine Person bekommt, sobald die Antwort läuft.",
              "Decode-Rate, tok/s pro Nutzer"),
             ("Gesamtleistung", "Alle Tokens pro Sekunde zusammen, für alle Personen.",
              "Durchsatz (Throughput)"),
             ("gleichzeitig", "Wie viele Fragen der Server im selben Moment bearbeitet.",
              "Concurrency / Parallelität"),
             ("„normal“ und „fast alle“", "Normal: Die Hälfte war schneller. Fast alle: 95 von 100 "
              "waren schneller.", "Median (p50) und p95"),
             ("Notizblock", "Hier merkt sich der Server die laufenden Gespräche. Ist er voll, "
              "müssen andere warten.", "KV-Cache"),
             ("Bruchpunkt", "Ab hier gibt es Fehler oder man wartet viel zu lange.", "Sättigung / Überlast")]
    for i, (t, e, tech) in enumerate(terms):
        c, r = divmod(i, 4)
        x, yy = 0.6 + c * 6.2, 1.65 + r * 1.32
        rect(sl, x, yy, 5.9, 1.2, CARD)
        text(sl, x + 0.25, yy + 0.1, 5.4, 0.35, t, size=15, bold=True, color=S1)
        text(sl, x + 0.25, yy + 0.42, 5.4, 0.5, e, size=12, color=INK)
        text(sl, x + 0.25, yy + 0.92, 5.4, 0.22, "Fachbegriff: " + tech, size=10, color=MUTED, italic=True)

    # ═══ 5 Einer allein ═══
    if base:
        sl = d.slide()
        title(sl, "Wie schnell ist er, wenn nur einer fragt?",
              f"Kurz gesagt: Er antwortet {speed_feel(base['user_tps_p50'])}.")
        e2e = base.get("e2e_p50")
        cards = [(fmt(base["user_tps_p50"], 1), "Tokens pro Sekunde",
                  f"≈ {n0(words(base['user_tps_p50']))} Wörter pro Sekunde. Zum Vergleich: "
                  "Ein Mensch liest etwa 4–5 Wörter pro Sekunde.", S1),
                 (secs(base["ttft_p50"]), "bis das erste Wort kommt",
                  cap(time_feel(base["ttft_p50"])) + "."
                  + (f" Er denkt erst nach – die eigentliche Antwort beginnt nach "
                     f"{secs(base['ttfc_p50'])}." if base.get("ttfc_p50") is not None else ""), INK),
                 (secs(e2e), "für eine ganze Antwort",
                  f"Eine Antwort war {n0(out_len)} Tokens lang (≈ {n0(words(out_len))} Wörter, "
                  "etwa ein Absatz).", INK)]
        for i, (v, l, m, col) in enumerate(cards):
            stat(sl, 0.6 + i * 4.1, 1.75, 3.85, 2.55, v, l, m, color=col, vsize=40)
        rect(sl, 0.6, 4.6, W - 1.2, 1.95, PANEL)
        text(sl, 0.95, 4.85, 11.5, 0.4, "Warum dieser Test?", size=15, bold=True, color=S1)
        pc = base.get("prefix_compare")
        text(sl, 0.95, 5.25, 11.5, 1.2,
             "Das ist der Bestwert: Der Erzähler hat nur ein Kind vor sich und kann sich ganz auf "
             "es konzentrieren. Alle anderen Tests zeigen, wie viel davon übrig bleibt, wenn mehr "
             f"los ist. Gemessen mit {base['n']} Fragen, jede ≈ {n0(words(in_len))} Wörter lang."
             + (f" Beginnen alle Fragen mit demselben Einleitungstext (≈ {n0(words(pc['prefix_tokens']))} "
                "Wörter), merkt er sich diesen: Das erste Wort kommt dann nach %s statt %s."
                % secs_pair(pc["ttft_shared"], pc["ttft_unique"]) if pc and pc.get("ttft_unique") else ""),
             size=13.5)
        footer(sl, f"Fachbegriffe: Decode-Rate (Median) {fmt(base['user_tps_p50'], 1)} tok/s · "
                   f"TTFT p50 {fmt(base['ttft_p50'], 3)} s / p95 {fmt(base['ttft_p95'], 3)} s · "
                   f"TPOT {fmt((base['tpot_p50'] or 0) * 1000, 1)} ms · Fehler {base['errors']}"
                   + (f" · Prefix-Cache-Treffer {fmt(base['prefix_hit_pct'], 0)} %"
                      if base.get("prefix_hit_pct") is not None else ""))

    # ═══ 6–8 Immer mehr gleichzeitig ═══
    if steps:
        cats = [s["concurrency"] for s in steps]
        tp = [s["out_tps"] for s in steps]
        # 6: Gesamtleistung
        sl = d.slide()
        gain = (top["out_tps"] / steps[0]["out_tps"]) if steps[0]["out_tps"] else None
        title(sl, "Wie viel schafft er insgesamt?",
              f"Kurz gesagt: Mit mehr Fragen gleichzeitig schafft er bis zu "
              f"{fmt(gain, 0 if gain and gain >= 10 else 1)}-mal so viel – dann ist Schluss."
              if gain else None)
        line_chart(sl, 0.5, 1.7, 8.2, 5.0, cats, [("Gesamtleistung", tp)], None,
                   "Tokens pro Sekunde (alle zusammen)", "Fragen gleichzeitig",
                   colors=(S1,), number_format="0")
        explain(sl,
                "Jeder Punkt ist eine Stufe: links wenige Fragen gleichzeitig, rechts viele. "
                "Je höher der Punkt, desto mehr Text schreibt der Server insgesamt pro Sekunde.",
                f"Anfangs steigt die Linie: Der Erzähler kann mehreren Kindern gleichzeitig "
                f"antworten. Ab etwa {sat or top['concurrency']} gleichzeitig wird die Linie flach "
                f"– mehr als ≈ {n0(top['out_tps'])} Tokens pro Sekunde schafft er nicht.")
        footer(sl, "Fachbegriff: Durchsatz (Output-Tokens/s), gemessen im eingeschwungenen Zustand "
                   "ohne Anlauf- und Auslaufphase.")

        # 7: Tempo pro Person
        sl = d.slide()
        ok_n = rec or 0
        title(sl, "Wie schnell ist es für jeden Einzelnen?",
              f"Kurz gesagt: Bis {ok_n} gleichzeitig bekommt jeder genug Tempo." if rec else
              "Kurz gesagt: Schon bei wenigen gleichzeitig wird es für jeden zäh.")
        line_chart(sl, 0.5, 1.7, 8.2, 5.0, cats,
                   [("Tempo pro Person", [s["user_tps_p50"] for s in steps]),
                    (f"Ziel: mindestens {n0(slo_u)}", [slo_u] * len(cats))], None,
                   "Tokens pro Sekunde für eine Person", "Fragen gleichzeitig",
                   colors=(S2, REF), dashed=(1,))
        last = steps[-1]
        stall_s = rec_step or steps[0]
        stall = stall_s.get("stall_p95")
        explain(sl,
                "Die orange Linie zeigt, wie schnell die Antwort für eine einzelne Person "
                "läuft. Die gestrichelte Linie ist unser Ziel. Solange Orange darüber liegt, "
                "ist alles flüssig.",
                f"Der Erzähler teilt seine Kraft auf alle auf. Allein bekommt man "
                f"{fmt(steps[0]['user_tps_p50'], 0)} Tokens/s, bei {last['concurrency']} gleichzeitig "
                f"nur noch {fmt(last['user_tps_p50'], 0)} – das ist {speed_feel(last['user_tps_p50'])}."
                + (f" Stocken: Bei {stall_s['concurrency']} gleichzeitig hängt die Antwort "
                   f"höchstens {secs(stall)} am Stück ({stall_feel(stall)})."
                   if stall is not None else ""))
        footer(sl, "Fachbegriffe: Decode-Rate pro Nutzer (Median), Ziel = SLO für die "
                   "Ausgabegeschwindigkeit; Stocken = längste Pause zwischen zwei Textstücken "
                   "einer Antwort (p95).")

        # 8: Wartezeit + Ampel
        sl = d.slide()
        title(sl, "Wie lange muss man warten, bis die Antwort beginnt?",
              f"Kurz gesagt: Ab {brk['concurrency']} gleichzeitig ist Schluss – {brk['reason']}."
              if brk else "Kurz gesagt: Kein Zusammenbruch im getesteten Bereich.")
        line_chart(sl, 0.5, 1.65, 8.2, 3.95, cats,
                   [("normal (die Hälfte wartet kürzer)", [s["ttft_p50"] for s in steps]),
                    ("fast alle (95 von 100 warten kürzer)", [s["ttft_p95"] for s in steps]),
                    (f"Ziel: höchstens {fmt(slo_t, 1)} s", [slo_t] * len(cats))], None,
                   "Sekunden bis zum ersten Wort", None, colors=(S1, S2, REF), dashed=(2,),
                   number_format="0.0")
        # Ampel-Leiste
        n = len(steps)
        box_w = min(1.0, 8.0 / n)
        x0 = 0.6 + (8.1 - box_w * n) / 2
        text(sl, 0.6, 5.7, 8.1, 0.3, "Ampel je Stufe (Zahl = Fragen gleichzeitig)", size=11,
             color=MUTED, align=PP_ALIGN.CENTER)
        for i, s in enumerate(steps):
            if meets_slo(s, slo_t, slo_u):
                col = GOOD
            elif s["error_rate"] > 0 or (brk and brk["concurrency"] == s["concurrency"]):
                col = BAD
            else:
                col = WARN
            dd = min(0.62, box_w - 0.12)
            dot(sl, x0 + i * box_w + (box_w - dd) / 2, 6.05, dd, col, s["concurrency"],
                size=11 if s["concurrency"] < 100 else 9)
        explain(sl,
                "Blau: wie lange man normalerweise wartet. Orange: wie lange die Pechvögel "
                "warten. Die Ampel unten fasst jede Stufe zusammen: Grün = gut, Gelb = zu "
                "langsam, Rot = Fehler oder überlastet.",
                "Solange der Erzähler noch Platz hat, fängt er sofort an. Ist sein Kopf voll, "
                "müssen neue Kinder in einer Schlange warten – dann steigt die Wartezeit steil an."
                + (f" Empfehlung: höchstens {rec} gleichzeitig." if rec else ""))
        footer(sl, "Fachbegriffe: TTFT p50 / p95; Ampel grün = SLO eingehalten (TTFT p95 und "
                   "Decode-Rate), rot = Fehler oder Bruchkriterium.")

    # ═══ 8b Echter Andrang ═══
    if rstages:
        sl = d.slide()
        r_ok = cap_.get("rate_ok")
        title(sl, "Was passiert, wenn die Fragen zufällig kommen?",
              (f"Kurz gesagt: Bis {per_min(r_ok)} Fragen pro Minute bleibt es flüssig – das reicht "
               f"für ≈ {approx(cap_.get('users_rate'))} aktive Nutzer.") if r_ok else
              "Kurz gesagt: Schon bei der kleinsten getesteten Menge wird es zu langsam.")
        cats = [per_min(s["rate"]) for s in rstages]
        line_chart(sl, 0.5, 1.7, 8.2, 5.0, cats,
                   [("normal (die Hälfte wartet kürzer)", [s["ttft_p50"] for s in rstages]),
                    ("fast alle (95 von 100 warten kürzer)", [s["ttft_p95"] for s in rstages]),
                    (f"Ziel: höchstens {fmt(slo_t, 1)} s", [slo_t] * len(cats))], None,
                   "Sekunden bis zum ersten Wort", "Fragen pro Minute (zufällig verteilt)",
                   colors=(S1, S2, REF), dashed=(2,), number_format="0.0")
        explain(sl,
                "Im Alltag kommen Fragen nicht im Takt, sondern zufällig: mal drei auf einmal, "
                "dann eine Weile keine. Genau so haben wir hier gefragt. Nach rechts kommen mehr "
                "Fragen pro Minute.",
                "Solange der Erzähler schneller antwortet, als neue Kinder kommen, bleibt die Linie "
                "flach. Kommen mehr, als er schafft, wird die Schlange immer länger – die Wartezeit "
                "schießt nach oben."
                + (f" Das passierte ab {per_min(rbrk['rate'])} Fragen pro Minute." if rbrk else "")
                + f" Annahme: Jede Person macht {n0(think)} s Pause zwischen zwei Fragen.")
        footer(sl, "Fachbegriffe: offene Last (Poisson-Ankünfte), Gesetz von Little. Ziele erfüllt "
                   "(Goodput): " + " · ".join(f"{per_min(s['rate'])}/min {fmt(s.get('goodput_pct'), 0)} %"
                                               for s in rstages))

    # ═══ 9 Lange Texte ═══
    if ctx and ctx.get("steps"):
        cs = [s for s in ctx["steps"] if s["single"]["ok"] > 0]
        sl = d.slide()
        lim = ctx.get("limit")
        title(sl, "Wie viel Text kann er auf einmal lesen?",
              f"Kurz gesagt: Bis ≈ {n0(pages(ctx['max_ok']))} Buchseiten klappt es"
              + (f", mehrere lange Texte gleichzeitig überfordern ihn ab ≈ {n0(pages(lim['input_len']))} Seiten."
                 if lim and lim["mode"] != "einzeln" else ".") if ctx.get("max_ok") else None)
        if cs:
            cc = cfg.get("context_concurrency", 1)
            series = [("einer allein", [s["single"]["ttft_p50"] for s in cs])]
            if any(s.get("parallel") for s in cs):
                series.append((f"{cc} gleichzeitig (fast alle)",
                               [s["parallel"]["ttft_p95"] if s.get("parallel") and s["parallel"]["ok"] else None
                                for s in cs]))
            line_chart(sl, 0.5, 1.7, 8.2, 5.0, [ctx_label(s["input_len"]) for s in cs], series, None,
                       "Sekunden bis zum ersten Wort", "Länge des Textes (Tokens, ≈ Buchseiten)",
                       number_format="0.0")
            pre = [s.get("prefill_tps") for s in cs if s.get("prefill_tps")]
            explain(sl,
                    "Nach rechts werden die Texte länger, die wir dem Server zum Lesen geben. "
                    "Nach oben: wie lange er liest, bevor er antwortet.",
                    f"Er liest etwa {n0(pages(max(pre)) if pre else None)} Buchseiten pro Sekunde. "
                    "Lange Texte brauchen viel Platz auf seinem Notizblock. Kommen mehrere davon "
                    "gleichzeitig, kann der Platz ausgehen – dann gibt es Fehler."
                    + (f" Das passierte bei ≈ {n0(pages(lim['input_len']))} Seiten ({lim['mode']})."
                       if lim else ""))
        footer(sl, "Fachbegriffe: Kontextlänge, Prefill-Geschwindigkeit, KV-Cache. "
                   "1 Buchseite ≈ 300 Wörter ≈ 400 Tokens (grobe Faustregel).")

    # ═══ 10 Ansturm ═══
    if spike:
        b = spike["burst"]
        sl = d.slide()
        ok_all = b["errors"] == 0
        title(sl, "Was passiert beim großen Ansturm?",
              f"Kurz gesagt: {spike['size']} Fragen auf einmal – "
              + ("alle wurden beantwortet" if ok_all else f"{b['errors']} gingen verloren")
              + (", danach war er wieder normal." if spike.get("recovered") else
                 ", danach war er nicht wieder normal."))
        ts = spike.get("ttft_sorted") or []
        if ts:
            step_ = max(1, len(ts) // 64)
            bar_chart(sl, 0.5, 1.7, 8.2, 5.0, list(range(1, len(ts) + 1, step_)), ts[::step_], None,
                      "Sekunden bis zum ersten Wort", "Fragen, sortiert: links die schnellste, "
                      "rechts die langsamste")
        explain(sl,
                "Jeder Balken ist eine Frage. Je höher der Balken, desto länger musste sie warten. "
                "Sie sind der Größe nach sortiert.",
                ("Siehst du Treppenstufen, nimmt der Erzähler immer eine Gruppe dran – die anderen "
                 "warten auf die nächste Runde. ") +
                f"Die längste Wartezeit war {secs(b['ttft_max'])} – {time_feel(b['ttft_max'])}.")
        footer(sl, "Fachbegriffe: Burst-/Spike-Test; Erholung = Einzelanfragen danach wieder auf "
                   "Baseline-Niveau (TTFT ≤ 1,5×, Decode-Rate ≥ 80 %).")

    # ═══ 11 Ausdauer ═══
    if soak and any(bk["s"] for bk in soak.get("buckets", [])):
        bks = [bk for bk in soak["buckets"] if bk["s"]]
        dg = (soak.get("degradation") or {}).get("tps_change_pct")
        sl = d.slide()
        tired = dg is not None and dg < -10
        dd = dg or 0
        change = ("genauso schnell wie am Anfang" if abs(dd) < 2 else
                  f"{fmt(abs(dd), 0)} % {'langsamer' if dd < 0 else 'schneller'} als am Anfang")
        title(sl, "Wird er mit der Zeit müde?",
              f"Kurz gesagt: {'Ja' if tired else 'Nein'} – nach {mins(soak['minutes'])} "
              f"Dauerbetrieb ist er {change}.")
        ch = xy_chart(sl, 0.5, 1.7, 8.2, 5.0,
                 [("Gesamtleistung", [(bk["t_min"], bk["s"]["out_tps"] or 0) for bk in bks])], None,
                 "Tokens pro Sekunde (alle zusammen)", colors=(S1,))
        ch.value_axis.minimum_scale = 0  # Null-Linie zeigen, sonst wirkt Rauschen wie Absturz
        tot = soak["total"]
        explain(sl,
                f"Wir haben den Server {mins(soak['minutes'])} lang mit "
                f"{soak['concurrency']} Fragen gleichzeitig beschäftigt. Die Linie zeigt seine "
                "Leistung Minute für Minute. Bleibt sie flach, ist alles gut.",
                ("Die Leistung sinkt mit der Zeit. Das deutet auf Hitze hin: Wird der Chip zu warm, "
                 "bremst er sich selbst. Kühlung und Aufstellort prüfen."
                 if tired else
                 "Die Leistung bleibt gleich. Der Server wird bei dieser Last nicht müde – kein "
                 "Hitzeproblem, kein Speicherleck.")
                + f" Insgesamt {n0(tot['ok'])} Antworten, {tot['errors']} Fehler.")
        footer(sl, "Fachbegriffe: Soak-Test, thermische Drosselung, Memory Leak.")

    # ═══ 11b Antwortqualität unter Last ═══
    if pr and pr.get("groups"):
        groups = pr["groups"]
        sl = d.slide()
        ref = pr.get("alone_pct")
        load_ = pr.get("load_pct")
        drop = (ref - load_) if ref is not None and load_ is not None else None
        title(sl, "Bleiben die Antworten richtig, wenn viel los ist?",
              ("Kurz gesagt: Ja – unter Last genauso gut wie allein." if drop is not None and drop <= 3 else
               f"Kurz gesagt: Nein – unter Last nur {fmt(load_, 0)} % richtig statt {fmt(ref, 0)} %."
               if drop is not None else None))
        ch = bars_chart(sl, 0.5, 1.7, 8.2, 5.0, [g["label"] for g in groups],
                        [("richtig beantwortet", [g["pct"] for g in groups])],
                        "Anteil richtiger Antworten (%)", colors=(S3,), label_format='0" %"')
        ch.value_axis.maximum_scale = 100
        kinds = ", ".join(f"{PROBE_KINDS.get(k, k)} {fmt(v, 0)} %" for k, v in pr.get("by_kind", {}).items())
        explain(sl,
                "Während der Lasttests haben wir nebenher einfache Kontrollfragen gestellt, deren "
                "Antwort feststeht: Rechnen, Sortieren, Hauptstädte und eine Antwort im festen "
                "Format (JSON). Jede Säule zeigt, wie viele davon richtig waren – links ohne Last.",
                ("Viel Andrang macht den Erzähler langsamer, aber nicht schlechter: Er antwortet "
                 "unter Last genauso richtig wie allein. "
                 if drop is not None and drop <= 3 else
                 "Unter Last macht der Erzähler mehr Fehler. Das deutet auf ein Problem mit den "
                 "Einstellungen hin (z. B. zu starke Kompression des Notizblocks) – prüfen. ")
                + f"Nach Art: {kinds}.")
        fails = sum(g["falsch"] for g in groups)
        cut = sum(g["abgeschnitten"] for g in groups)
        err = sum(g["fehler"] for g in groups)
        footer(sl, f"{pr['n_alone'] + pr['n_load']} Kontrollfragen · Temperatur 0 · falsch {fails} · "
                   f"abgeschnitten {cut} · Fehler {err}. Kontrollfragen laufen nebenher, höchstens "
                   "eine zusätzlich zur Last.")
        if pr.get("wrong"):
            notes(sl, "Beispiele für nicht richtige Antworten:\n" + "\n".join(
                f"[{w['label']}] {w['frage']} → {w['antwort']!r} ({w['status']})" for w in pr["wrong"]))

    # ═══ 11c Wirtschaftlichkeit ═══
    if eco:
        sl = d.slide()
        be = eco.get("breakeven_pct")
        title(sl, "Was kostet uns das?",
              (f"Kurz gesagt: Ab {fmt(be, 0)} % Auslastung ist der eigene Server günstiger als die Cloud."
               if be is not None and be <= 100 else
               "Kurz gesagt: Bei den eingetragenen Preisen wäre die Cloud günstiger." if be is not None else None))
        levels = (0.25, 0.5, 0.75, 1.0)
        bars_chart(sl, 0.5, 1.7, 8.2, 5.0, [f"{int(f * 100)} % ausgelastet" for f in levels],
                   [("eigener Server", [eco["own_month"]] * len(levels)),
                    ("Cloud, gleiche Menge", [eco["cloud_full_month"] * f for f in levels])],
                   "Euro pro Monat", colors=(S1, S2), label_format='#,##0 "€"')
        explain(sl,
                "Der eigene Server kostet jeden Monat gleich viel, egal wie viel er arbeitet "
                "(Abschreibung und Strom). Die Cloud kostet pro Wort. Die Säulen zeigen beides "
                "für unterschiedlich starke Nutzung während der Arbeitszeit.",
                f"Bei voller Nutzung beantwortet er ≈ {n0(eco['answers_month'])} Fragen im Monat. "
                f"1.000 Antworten kosten dann {eur(eco['own_per_1000'], 2)} statt "
                f"{eur(eco['cloud_per_1000'], 2)} in der Cloud"
                + (f", pro aktivem Nutzer {eur(eco['own_per_user'], 2)} im Monat." if eco.get("own_per_user") else ".")
                + " Nicht eingerechnet: Wartung, Einrichtung und der Vorteil, dass keine Daten das Haus verlassen.")
        footer(sl, f"Annahmen: Kaufpreis {eur(eco['hw_price'])} über {n0(eco['amort_months'])} Monate · "
                   f"{n0(eco['server_watts'])} W rund um die Uhr zu {fmt(eco['power_price'], 2)} €/kWh · "
                   f"Cloud {fmt(eco['cloud_price_in'], 2)} € / {fmt(eco['cloud_price_out'], 2)} € je Mio. "
                   f"Tokens rein/raus · {n0(eco['work_hours'])} Arbeitsstunden/Monat · Ø {n0(eco['in_avg'])} "
                   f"Tokens rein, {n0(eco['out_avg'])} raus je Antwort.")

    # ═══ 12 Maschinenraum ═══
    mon = summary.get("monitor") or []

    def ser(key, scale=1.0):
        st = max(1, len(mon) // 300)
        return [(round(m["t"] / 60, 2), m[key] * scale) for m in mon[::st] if m.get(key) is not None]

    panels = []
    if ser("kv"):
        panels.append(([("belegt", ser("kv"))], "Wie voll ist sein Notizblock? (%)",
                       "KV-Cache-Belegung", (S3,)))
    if ser("waiting"):
        panels.append(([("wird bedient", ser("running")), ("wartet in der Schlange", ser("waiting"))],
                       "Wie viele werden bedient – wie viele warten?", "running / waiting", (S1, S2)))
    if ser("gpu_temp"):
        panels.append(([("Temperatur", ser("gpu_temp"))], "Wie warm ist der Chip? (°C)",
                       "GPU-Temperatur", (S2,)))
    if ser("gpu_clock"):
        panels.append(([("Takt", ser("gpu_clock"))], "Wie schnell arbeitet der Chip? (MHz)",
                       "GPU-Takt", (S4,)))
    if ser("ram_used_gb"):
        panels.append(([("belegt", ser("ram_used_gb"))], "Wie viel Speicher ist belegt? (GB)",
                       "RAM", (S1,)))
    if panels:
        sl = d.slide()
        title(sl, "Ein Blick in den Maschinenraum",
              "Was im Server passierte, während wir getestet haben")
        panels = panels[:4]
        pos = [(0.5, 1.6), (6.85, 1.6), (0.5, 4.2), (6.85, 4.2)] if len(panels) > 2 else \
              [(0.5, 1.6), (6.85, 1.6)]
        hh = 2.5 if len(panels) > 2 else 4.6
        for (series, t, _tech, cols), (px, py) in zip(panels, pos):
            ch = xy_chart(sl, px, py, 6.0, hh, series, t, None, "Minuten seit Start", colors=cols)
        footer(sl, "Der Notizblock ist der KV-Cache: Dort merkt sich der Server die laufenden "
                   "Gespräche. Ist er voll, müssen neue Fragen in der Schlange warten.")
        notes(sl, "Quelle: Prometheus-/metrics des Modellservers" +
              (" und nvidia-smi / /proc/meminfo lokal." if meta.get("local_monitor") else "."))

    # ═══ 13 Datenqualität ═══
    items = quality.get("items", [])
    if items:
        sl = d.slide()
        title(sl, "Wie sehr können wir den Zahlen trauen?",
              "Wir haben nicht nur den Server geprüft, sondern auch die Messung selbst.")
        rect(sl, 9.6, 0.45, 3.2, 0.95, CARD)
        dot(sl, 9.85, 0.74, 0.36, q_color)
        text(sl, 10.35, 0.52, 2.4, 0.3, "Verlässlichkeit", size=11, color=MUTED)
        text(sl, 10.35, 0.8, 2.4, 0.45, quality["overall"].upper(), size=20, bold=True, color=q_color)
        cols = 2 if len(items) > 5 else 1
        per = (len(items) + cols - 1) // cols
        cw = (W - 1.2 - 0.3 * (cols - 1)) / cols
        ih = min(1.05, 5.25 / per)
        for i, it in enumerate(items):
            c, r = divmod(i, per)
            x, yy = 0.6 + c * (cw + 0.3), 1.7 + r * ih
            dot(sl, x, yy + 0.06, 0.26, STATUS[it["status"]])
            text(sl, x + 0.42, yy, cw - 0.45, 0.3,
                 [[(it["title"], {"bold": True}), ("   " + it["value"], {"color": MUTED})]], size=13)
            text(sl, x + 0.42, yy + 0.32, cw - 0.45, ih - 0.34, it["text"], size=11, color=INK2)
        footer(sl, "Grün = in Ordnung · Gelb = mit Vorsicht lesen · Rot = Ergebnis verfälscht · "
                   "Grau = Hinweis")

    # ═══ 14 Fazit ═══
    sl = d.slide(dark=True)
    text(sl, 0.6, 0.5, 12, 0.7, "Was heißt das für uns?", size=32, bold=True, color="FFFFFF",
         font="Cambria")
    f = []
    if base:
        f.append(("Einer allein:", f"flott – {speed_feel(base['user_tps_p50'])}, erstes Wort nach "
                                   f"{secs(base['ttft_p50'])}."))
    if steps or rstages:
        f.append(("Gut nutzbar:", (f"für ≈ {approx(users)} aktive Nutzer"
                                   + (f" ({rec} Fragen gleichzeitig)." if rec else "."))
                  if users else
                  "keine Stufe erreicht die Ziele – Ziele, Modell oder Einstellungen prüfen."))
    if steps:
        f.append(("Grenze:", f"ab {brk['concurrency']} gleichzeitig geht es kaputt ({brk['reason']})."
                  if brk else f"ab {saturated} gleichzeitig steigt die Leistung nicht mehr."
                  if saturated else f"bis {steps[-1]['concurrency']} gleichzeitig nicht kaputt gegangen."))
    if rec_step and rec_step.get("stall_p95") is not None:
        f.append(("Flüssig:", f"Antworten hängen höchstens {secs(rec_step['stall_p95'])} am Stück – "
                              f"{stall_feel(rec_step['stall_p95'])}."))
    if ctx and ctx.get("max_ok"):
        f.append(("Lange Texte:", f"bis ≈ {n0(pages(ctx['max_ok']))} Buchseiten."
                  + (f" Mehrere lange Texte gleichzeitig: Probleme ab ≈ {n0(pages(ctx['limit']['input_len']))} Seiten."
                     if ctx.get("limit") and ctx["limit"]["mode"] != "einzeln" else "")))
    if spike:
        f.append(("Ansturm:", "gut überstanden." if spike.get("recovered") and spike["burst"]["errors"] == 0
                  else "nicht sauber überstanden."))
    if soak and soak.get("degradation") and soak["degradation"].get("tps_change_pct") is not None:
        dg = soak["degradation"]["tps_change_pct"]
        f.append(("Ausdauer:", "wird nicht müde." if dg > -10 else "wird mit der Zeit langsamer."))
    if pr and pr.get("load_pct") is not None:
        f.append(("Antworten:", f"unter Last {fmt(pr['load_pct'], 0)} % richtig"
                  + (f" (allein {fmt(pr['alone_pct'], 0)} %)." if pr.get("alone_pct") is not None else ".")))
    if eco and eco.get("breakeven_pct") is not None:
        f.append(("Kosten:", f"{eur(eco['own_month'])} im Monat – "
                  + (f"günstiger als die Cloud ab {fmt(eco['breakeven_pct'], 0)} % Auslastung."
                     if eco["breakeven_pct"] <= 100 else "die Cloud wäre günstiger.")))
    f.append(("Messung:", f"Verlässlichkeit {quality['overall']}."))
    text(sl, 0.6, 1.55, 7.4, 5.2, [[(h + " ", {"bold": True, "color": "FFFFFF"}), (b, {"color": "CBD5E1"})]
                                   for h, b in f], size=15, space_after=10)
    recs = []
    if rec:
        rng = f"{rec}–{sat}" if sat and sat > rec else f"{rec}"
        recs.append(("Anzahl begrenzen", f"Den Server auf höchstens {rng} gleichzeitige Fragen "
                     "einstellen. Wer später kommt, wartet kurz statt alle auszubremsen.",
                     "vLLM: --max-num-seqs"))
    preempted = [s for s in steps if s.get("preemptions") and (not rec or s["concurrency"] <= rec)]
    if preempted:
        recs.append(("Notizblock vergrößern", f"Schon bei {preempted[0]['concurrency']} gleichzeitig "
                     "musste der Server angefangene Antworten verwerfen, weil der Notizblock voll war.",
                     "Preemption: --gpu-memory-utilization, KV-Cache fp8"))
    if ctx and ctx.get("limit") and ctx["limit"]["mode"] != "einzeln":
        recs.append(("Lange Texte begrenzen", "Die maximale Textlänge senken oder den Notizblock "
                     "sparsamer machen.", "--max-model-len, KV-Cache fp8"))
    if soak and soak.get("degradation") and (soak["degradation"].get("tps_change_pct") or 0) < -10:
        recs.append(("Kühlung prüfen", "Aufstellort und Belüftung des Servers prüfen.",
                     "thermische Drosselung"))
    if pr and pr.get("load_pct") is not None and (pr.get("alone_pct") or 100) - pr["load_pct"] > 3:
        recs.append(("Antwortqualität prüfen", "Unter Last gab es mehr falsche Antworten. Einstellungen "
                     "zur Kompression und zum Notizblock prüfen.", "Quantisierung, KV-Cache-dtype"))
    if quality["overall"] != "hoch":
        recs.append(("Test wiederholen", "Den vollständigen Test laufen lassen, wenn niemand "
                     "sonst den Server nutzt.", "siehe Folie Datenqualität"))
    recs.append(("Nach Änderungen neu testen", "Nach jeder Einstellungsänderung erneut messen "
                 "und die Berichte vergleichen.", "Regressionstest"))
    rect(sl, 8.35, 1.5, 4.4, 5.3, "1E293B")
    text(sl, 8.65, 1.72, 3.9, 0.4, "Nächste Schritte", size=17, bold=True, color="FFFFFF")
    yy = 2.25
    for h, b_, tech in recs[:4]:
        text(sl, 8.65, yy, 3.85, 1.1, [[("→ " + h, {"bold": True, "color": "FFFFFF"})],
                                       [(b_, {"color": "CBD5E1"})],
                                       [(tech, {"color": "94A3B8", "italic": True, "size": 10})]],
             size=12)
        yy += 1.13

    # ═══ 15 Anhang: Zahlen ═══
    if steps:
        sl = d.slide()
        title(sl, "Anhang: alle Zahlen für Technik-Interessierte",
              "Stufen im eingeschwungenen Zustand · p50 = Median · p95 = 95. Perzentil")
        rows = [["Gleichzeitig", "Durchsatz tok/s", "tok/s pro Nutzer", "TTFT p50", "TTFT p95",
                 "Anfragen/s", "Ziele erfüllt", "Stocken p95", "Verdrängt", "Messungen", "Fehler",
                 "Schwankung", "Verzug Messr."]]
        for s in steps:
            col = GOOD if meets_slo(s, slo_t, slo_u) else (
                BAD if s["error_rate"] > 0 or (brk and brk["concurrency"] == s["concurrency"]) else WARN)
            rows.append([(f"{s['concurrency']}{' *' if s.get('fine') else ''}", col),
                         n0(s["out_tps"]), fmt(s["user_tps_p50"], 1), secs(s["ttft_p50"], 2),
                         secs(s["ttft_p95"], 2), fmt(s["req_s"], 2),
                         fmt(s.get("goodput_pct"), 0, " %"), secs(s.get("stall_p95"), 2),
                         fmt(s.get("preemptions")), str(s.get("n_window", s["n"])),
                         f"{s['error_rate'] * 100:.0f} %", fmt(s.get("stability_pct"), 0, " %"),
                         fmt(s.get("client_lag_ms"), 0, " ms")])
        rh = min(0.4, 5.0 / len(rows))
        table(sl, 0.6, 1.7, W - 1.2, rows,
              col_w=[1.0, 1.0, 0.95, 0.85, 0.85, 0.85, 0.9, 0.95, 0.85, 0.95, 0.75, 1.0, 1.233],
              font=10, row_h=rh)
        footer(sl, "* = Zwischenstufe aus der Feinsuche. Ziele erfüllt = Anteil der einzelnen Fragen, "
                   "die beide Ziele einhalten (Goodput). Stocken = längste Pause in einer Antwort. "
                   "Verdrängt = Antworten, die der Server wegen vollem KV-Cache neu beginnen musste.")

    if rstages:
        sl = d.slide()
        title(sl, "Anhang: Echter Andrang – alle Zahlen",
              "Zufällige Ankünfte (Poisson) · gemessen ab Ende der Anlaufzeit")
        rows = [["Fragen/min", "tatsächlich", "TTFT p50", "TTFT p95", "tok/s pro Nutzer",
                 "Durchsatz tok/s", "Ziele erfüllt", "Stocken p95", "max. gleichz.", "Schlange ×",
                 "Verdrängt", "Fehler"]]
        for s in rstages:
            col = GOOD if rate_ok(s, slo_t, slo_u) else BAD if s.get("overloaded") or s["error_rate"] > 0 else WARN
            rows.append([(per_min(s["rate"]), col), per_min(s.get("rate_actual")),
                         secs(s["ttft_p50"], 2), secs(s["ttft_p95"], 2), fmt(s["user_tps_p50"], 1),
                         n0(s["out_tps"]), fmt(s.get("goodput_pct"), 0, " %"),
                         secs(s.get("stall_p95"), 2), fmt(s.get("max_inflight")),
                         fmt(s.get("queue_growth"), 1), fmt(s.get("preemptions")),
                         f"{s['error_rate'] * 100:.0f} %"])
        rh = min(0.4, 5.0 / len(rows))
        table(sl, 0.6, 1.7, W - 1.2, rows,
              col_w=[1.1, 1.0, 0.95, 0.95, 1.1, 1.1, 1.0, 1.0, 1.0, 0.95, 0.9, 1.033],
              font=10.5, row_h=rh)
        footer(sl, "Schlange × = Wartezeit der letzten Ankünfte geteilt durch die der ersten "
                   "(> 2 = die Schlange wächst). max. gleichz. = höchstens so viele Fragen waren "
                   "gleichzeitig offen.")

    # Fehler-Anhang
    agg = Counter()
    for key, blk in (("Einer allein", base), ("Ansturm", spike and spike["burst"]),
                     ("Ausdauer", soak and soak["total"]), ("Kontrolle", ctl and ctl["s"])):
        if blk:
            for e, n in blk["error_types"].items():
                agg[(key, e)] += n
    for s in steps:
        for e, n in s["error_types"].items():
            agg[(f"{s['concurrency']} gleichzeitig", e)] += n
    for s in rstages:
        for e, n in s["error_types"].items():
            agg[(f"{per_min(s['rate'])} Fragen/min", e)] += n
    for s in (ctx or {}).get("steps", []):
        for part in ("single", "parallel"):
            if s.get(part):
                for e, n in s[part]["error_types"].items():
                    agg[(f"Text ≈{n0(pages(s['input_len']))} Seiten", e)] += n
    if agg:
        sl = d.slide()
        explain_err = {"verbindung": "Server nicht erreichbar", "timeout": "keine Antwort in der Zeit",
                       "leere_antwort": "Antwort war leer",
                       "zu_viele_offen": "nicht gesendet – schon zu viele Fragen offen"}
        title(sl, "Anhang: aufgetretene Fehler")
        rows = [["Wo", "Was", "Bedeutung", "Anzahl"]]
        for (t, e), n in list(agg.items())[:13]:
            meaning = explain_err.get(e) or ("vom Server abgelehnt (Anfrage ungültig/zu lang)"
                                             if e.startswith("http_4") else
                                             "Serverfehler (oft: Speicher voll)" if e.startswith("http_5")
                                             else "anderer Fehler")
            rows.append([t, e, meaning, str(n)])
        table(sl, 0.6, 1.5, 11.0, rows, col_w=[3.0, 1.8, 5.0, 1.2], font=12, row_h=0.36)

    # ═══ Schluss: Zusammenfassung auf einer Folie ═══
    # Alle Ergebnisse als Ampel-Liste plus die eine Zahl, die man sich merken soll.

    rows = result_rows(summary)

    sl = d.slide()
    started = (meta.get("started") or "")[:10]
    title(sl, "Zusammenfassung", f"{cfg.get('model') or '?'} · getestet am {started}")

    # Links: die eine Zahl
    rect(sl, 0.6, 1.65, 3.9, 5.1, DARK_BG)
    text(sl, 0.9, 1.95, 3.3, 0.35, "REICHT FÜR", size=13, bold=True, color="94A3B8")
    text(sl, 0.9, 2.35, 3.3, 1.1, f"≈ {approx(users)}" if users else "–", size=54, bold=True,
         color="FFFFFF", font="Cambria")
    text(sl, 0.9, 3.5, 3.3, 0.5, "aktive Nutzer gleichzeitig", size=17, bold=True, color="FFFFFF")
    text(sl, 0.9, 4.05, 3.3, 1.3,
         (f"Gebraucht werden {target}. " if target else "")
         + f"Jede Person stellt etwa alle {n0(think)} s eine Frage (Lesen, Nachdenken, Tippen). "
         f"Ziel: erstes Wort nach ≤ {fmt(slo_t, 1)} s, Tempo ≥ {n0(slo_u)} Tokens/s."
         if users else "Keine getestete Last erfüllt die Ziele.",
         size=12, color="CBD5E1")
    dot(sl, 0.9, 5.95, 0.3, q_color)
    text(sl, 1.35, 5.9, 3.0, 0.45, f"Verlässlichkeit: {quality['overall']}", size=14, bold=True,
         color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)

    # Rechts: Ampel-Liste
    rows = rows[:14]
    rh = min(0.5, 5.1 / max(1, len(rows)))
    for i, (col, area, res) in enumerate(rows):
        yy = 1.65 + i * rh
        if i % 2 == 0:
            rect(sl, 4.8, yy, W - 5.4, rh, "F8FAFC", rounded=False)
        dot(sl, 5.0, yy + (rh - 0.22) / 2, 0.22, col or INFO)
        fs = 13 if rh >= 0.42 else 11.5
        text(sl, 5.4, yy, 2.75, rh, area, size=fs, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, 8.2, yy, W - 8.8, rh, res, size=fs - 0.5, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
    footer(sl, "Grün = in Ordnung · Gelb = beobachten · Rot = Handlungsbedarf · Grau = Hinweis. "
               "Nutzerzahl nach dem Gesetz von Little – eine Schätzung, keine Garantie.")

    d.save(path)
