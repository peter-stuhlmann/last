# -*- coding: utf-8 -*-
"""Berichte zum Ablauf-Lasttest: Schritte, Bildschirmfotos und aufgezeichnete Anfragen statt Seitengewicht."""
import os
import struct
import sys
import tempfile
import unittest
import zlib

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report_web as rw  # noqa: E402
from test_report_web import all_text, page, sample  # noqa: E402

STEPS = """öffne https://chat.example/login
tippe in "E-Mail" {email}
tippe in "Passwort" {passwort}
klicke "Anmelden"
warte auf Text "Willkommen"
messen ab hier
tippe in "Nachricht" "Was ist ein Lasttest?"
drücke Enter
warte bis ".antwort" fertig ist (2 s ruhig)
"""


def png(path, w=4, h=3):
    """Ein winziges, gültiges PNG – reicht, damit python-pptx es einbetten kann."""
    raw = b"".join(b"\x00" + b"\x7f\x7f\x7f" * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xffffffff)

    with open(path, "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def flow_sample(outdir=None):
    s = sample()
    s["config"] = {"kind": "flow", "base_url": "https://chat.example/login", "tests": "baseline,ramp,rate,spike",
                   "levels": "1,8,16,32", "level_seconds": 30, "think_time": 60.0, "target_users": 200,
                   "slo_load": 10.0, "spike_size": 16, "include_assets": False, "steps_text": STEPS,
                   "placeholders": ["email", "passwort"]}
    steps = []
    for i, line in enumerate([ln for ln in STEPS.splitlines() if ln.strip()], start=1):
        shot = f"schritt_{i:02d}.png"
        if outdir:
            png(os.path.join(outdir, shot))
        steps.append({"line": i, "text": line, "screenshot": shot, "ms": 300.0})
    s["prechecks"] = {"status": 200, "bytes": 1800, "assets": 2, "total_bytes": 2000,
                      "final_url": "https://chat.example/chat", "rtt_ms": 900.0, "busy": None, "files": [],
                      "requests": [{"method": "POST", "host": "chat.example", "path": "/api/conversations", "query": "",
                                    "status": 201, "stream": False, "body_bytes": 2, "header_names": ["content-type"],
                                    "dynamic": []},
                                   {"method": "POST", "host": "chat.example", "path": "/api/conversations/{{v1}}/messages",
                                    "query": "", "status": 200, "stream": True, "body_bytes": 32,
                                    "header_names": ["content-type"], "dynamic": ["v1"]}],
                      "main": 1, "steps": steps, "measured_ms": 4200.0}
    s["baseline"] = page(ttft_p50=0.9, ttft_p95=1.2, e2e_p50=4.0, e2e_p95=5.5, e2e_mean=4.2, req_s=0.25,
                         bytes_avg=1800, files_avg=0)
    del s["soak"]
    return s


class AblaufBericht(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.path = os.path.join(self.dir, "bericht.pptx")

    def test_begriffe_und_schritte(self):
        s = flow_sample(self.dir)
        rw.build_report(s, self.path)
        text = all_text(self.path)
        for expected in ("Wie viele Nutzer hält unsere Anwendung aus?", "Nutzer gleichzeitig",
                         "Einmal im echten Browser aufgezeichnet", "ab hier wird gemessen",
                         "POST /api/conversations/{v1}/messages", "Was der Browser gesehen hat",
                         "Zeile 7", "Antwortzeit ohne Andrang", "Anfragen je Durchlauf"):
            self.assertIn(expected, text)
        for forbidden in ("Besucher", "Website", "Gewicht der Seite", "Token", "Seitenaufrufe", "Anfragen pro Minute"):
            self.assertNotIn(forbidden, text)
        self.assertNotRegex(text, r"(None|nan)")
        # Eingebettet: Anfang, Stand vor „messen“ und die drei gemessenen Schritte
        pics = sum(1 for sl in Presentation(self.path).slides for sh in sl.shapes if sh.shape_type == 13)
        self.assertEqual(pics, 5)
        self.assertNotIn("Zeile 3", text)

    def test_ohne_fotos_trotzdem_bericht(self):
        s = flow_sample()               # keine PNGs im Ordner
        rw.build_report(s, self.path)
        text = all_text(self.path)
        self.assertIn("ab hier wird gemessen", text)
        self.assertNotIn("Was der Browser gesehen hat", text)

    def test_pdf_und_vergleich(self):
        pdf = os.path.join(self.dir, "bericht.pdf")
        rw.build_pdf(flow_sample(), pdf)
        with open(pdf, "rb") as f:
            self.assertTrue(f.read().startswith(b"%PDF"))
        before, after = flow_sample(), flow_sample()
        after["meta"]["started"] = "2026-10-06T09:00:00"
        after["baseline"] = page(e2e_p50=2.0, ttft_p50=0.4)
        cmp_path = os.path.join(self.dir, "vergleich.pptx")
        rw.build_compare([before, after], cmp_path, ["vorher", "nachher"])
        text = all_text(cmp_path)
        self.assertIn("Ist die Anwendung schneller geworden?", text)
        self.assertNotIn("Besucher", text)


if __name__ == "__main__":
    unittest.main()
